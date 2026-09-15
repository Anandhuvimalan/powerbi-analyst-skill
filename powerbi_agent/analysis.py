"""Exact profiling within an explicit row budget; source values stay in memory."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from .core import BuildError, identifier


def words(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def is_key(name: str) -> bool:
    return bool(re.search(r"(?:Id|ID|Key)$|(?:^|_)(?:id|key)$", name)) or words(name) in {
        "id", "orderid", "customerid", "productid", "employeeid", "campaignid", "transactionid"}


def date_value(value):
    if isinstance(value, dt.datetime):
        if value.tzinfo or value.time() != dt.time():
            return None  # Keep timestamps intact; agent may explicitly derive a date key.
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        try:
            return dt.date.fromisoformat(value.strip())
        except ValueError:
            pass
    return None


def clean(value):
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, float) and not math.isfinite(value):
        raise BuildError("Non-finite numeric source value. Correct NaN/Infinity before modeling.")
    return value


@dataclass
class Dataset:
    name: str
    source: dict
    rows: list[dict]
    profile: dict


def profile(name: str, source: dict, rows: list[dict]) -> Dataset:
    if not rows:
        raise BuildError(f"{name}: source has no data rows.")
    names = list(rows[0])
    if not names or any(not n or not isinstance(n, str) for n in names):
        raise BuildError(f"{name}: nonempty string column headers are required.")
    if len({n.casefold() for n in names}) != len(names):
        raise BuildError(f"{name}: duplicate column names (Power BI is case insensitive).")
    cleaned = [{n: clean(row.get(n)) for n in names} for row in rows]
    columns, issues = [], []
    for col in names:
        vals = [row[col] for row in cleaned if row[col] is not None]
        if any(isinstance(v, (dict, list)) for v in vals):
            raise BuildError(f"{name}[{col}]: nested data requires an explicit flattening plan.")
        trim = any(isinstance(r.get(col), str) and r[col] != r[col].strip() for r in rows)
        kind = "string"
        if vals and all(isinstance(v, bool) for v in vals):
            kind = "boolean"
        elif vals and all(date_value(v) is not None for v in vals):
            kind = "dateTime"
            for row in cleaned:
                if row[col] is not None:
                    row[col] = date_value(row[col]).isoformat()
        elif vals and not is_key(col):
            try:
                numbers = [float(v) for v in vals]
                if all(math.isfinite(v) for v in numbers) and not any(isinstance(v, bool) for v in vals):
                    kind = "int64" if all(v.is_integer() and abs(v) < 2**53 for v in numbers) else "double"
                    for row in cleaned:
                        if row[col] is not None:
                            row[col] = int(float(row[col])) if kind == "int64" else float(row[col])
            except (ValueError, TypeError):
                pass
        if kind == "string":
            for row in cleaned:
                if row[col] is not None:
                    row[col] = str(row[col])
        vals = [r[col] for r in cleaned if r[col] is not None]
        distinct = len(set(vals))
        nulls = len(rows) - len(vals)
        role = "key" if is_key(col) else "date" if kind == "dateTime" else "numeric" if kind in {"double", "int64"} else "category"
        if role == "category" and distinct > max(50, len(rows) * .5):
            role = "description"
        columns.append({"name": col, "data_type": kind, "role": role, "distinct": distinct,
                        "nulls": nulls, "null_fraction": nulls / len(rows),
                        "unique": distinct == len(rows) and not nulls, "trim": trim,
                        "minimum": min(vals) if vals and kind in {"double", "int64", "dateTime"} else None,
                        "maximum": max(vals) if vals and kind in {"double", "int64", "dateTime"} else None,
                        "treatment": "Trim observed whitespace; keep nulls" if trim else "Type conversion only; keep nulls"})
        if nulls:
            issues.append({"column": col, "issue": "missing_values", "count": nulls, "action": "preserved"})
        if role == "key" and nulls:
            issues.append({"column": col, "issue": "null_key", "action": "do not use as dimension key"})
    duplicates = len(rows) - len({tuple(r[n] for n in names) for r in cleaned})
    if duplicates:
        issues.append({"issue": "duplicate_rows", "count": duplicates, "action": "preserved; grain requires business review"})
    return Dataset(name, source, cleaned, {"name": name, "row_count": len(rows), "complete": True,
        "columns": columns, "duplicate_rows": duplicates, "data_quality_issues": issues})


def discover(sources: list, base: Path, max_rows: int = 100000) -> list[Dataset]:
    datasets = []
    pending = list(sources)
    while pending:
        spec = pending.pop(0)
        spec = {"path": spec} if isinstance(spec, str) else dict(spec)
        if spec.get("kind") in {"sqlserver", "sqlite"}:
            datasets.append(database_source(spec, base, max_rows))
            continue
        path = (base / spec["path"]).resolve()
        if path.is_dir():
            pending.extend({"path": str(p)} for p in sorted(path.iterdir()) if p.suffix.lower() in {".csv", ".tsv", ".xlsx", ".json"})
            continue
        if not path.is_file():
            raise BuildError(f"Source does not exist: {path}")
        source = {**spec, "path": str(path), "kind": path.suffix.lower()[1:]}
        name = identifier(spec.get("name", path.stem))
        if path.suffix.lower() in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream, delimiter=spec.get("delimiter", "\t" if path.suffix == ".tsv" else ","))
                if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                    raise BuildError(f"{path}: duplicate or missing headers")
                rows = []
                for row in reader:
                    if None in row:
                        raise BuildError(f"{path}: a row has more values than headers")
                    rows.append(row)
                    if len(rows) > max_rows:
                        raise BuildError(f"{path}: exceeds max_rows={max_rows}; increase the budget or supply a curated source.")
            source["delimiter"] = reader.dialect.delimiter if isinstance(reader.dialect, csv.Dialect) else spec.get("delimiter", "\t" if path.suffix == ".tsv" else ",")
            datasets.append(profile(name, source, rows))
        elif path.suffix.lower() == ".json":
            if path.stat().st_size > 100_000_000:
                raise BuildError("JSON exceeds 100 MB; use CSV or a database source.")
            rows = json.loads(path.read_text(encoding="utf-8-sig"))
            if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
                raise BuildError("JSON sources must be an array of records.")
            if len(rows) > max_rows:
                raise BuildError(f"{path}: exceeds max_rows={max_rows}")
            headers = list(dict.fromkeys(k for r in rows for k in r))
            datasets.append(profile(name, source, [{k: r.get(k) for k in headers} for r in rows]))
        elif path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                sheets = [spec["sheet"]] if spec.get("sheet") else book.sheetnames
                for sheet in sheets:
                    iterator = book[sheet].iter_rows(values_only=True)
                    headers = next(iterator, ())
                    if not headers:
                        continue
                    if len(set(headers)) != len(headers) or any(not isinstance(h, str) or not h for h in headers):
                        raise BuildError(f"{sheet}: duplicate, blank or non-text Excel headers")
                    rows = []
                    for row in iterator:
                        if not any(v is not None for v in row):
                            continue
                        rows.append(dict(zip(headers, row)))
                        if len(rows) > max_rows:
                            raise BuildError(f"{sheet}: exceeds max_rows={max_rows}")
                    if rows:
                        datasets.append(profile(name if len(sheets) == 1 else identifier(sheet), {**source, "sheet": sheet}, rows))
            finally:
                book.close()
        else:
            raise BuildError(f"Unsupported source {path.suffix}. Convert legacy XLS to XLSX; supported: CSV, TSV, JSON, XLSX, SQLite, SQL Server.")
    if not datasets:
        raise BuildError("No tabular data found.")
    names = [d.name.casefold() for d in datasets]
    if len(set(names)) != len(names):
        raise BuildError("Sources produce duplicate table names. Supply a unique 'name' per source.")
    return datasets


def database_source(spec: dict, base: Path, max_rows: int) -> Dataset:
    table = spec.get("table", "")
    schema = spec.get("schema", "dbo")
    if not re.fullmatch(r"[\w ]+", table) or not re.fullmatch(r"\w+", schema):
        raise BuildError("Database sources require a table and a simple schema identifier; arbitrary SQL is not accepted.")
    if spec["kind"] == "sqlite":
        path = (base / spec["path"]).resolve()
        conn = sqlite3.connect(path.as_uri() + "?mode=ro", uri=True)
        query = f'SELECT * FROM "{table}" LIMIT {max_rows + 1}'
        spec = {**spec, "path": str(path)}
    else:
        try:
            import pyodbc
        except ImportError as exc:
            raise BuildError("SQL Server profiling requires pip install '.[sql]' and a configured ODBC driver.") from exc
        secret = os.environ.get(spec.get("connection_env", ""))
        if not secret:
            raise BuildError("SQL Server connection_env must name an environment variable containing an ODBC connection string.")
        conn = pyodbc.connect(secret, timeout=15, readonly=True)
        query = f"SELECT TOP ({max_rows + 1}) * FROM [{schema}].[{table}]"
    try:
        cursor = conn.execute(query)
        headers = [c[0] for c in cursor.description]
        rows = [dict(zip(headers, row)) for row in cursor.fetchall()]
    finally:
        conn.close()
    if len(rows) > max_rows:
        raise BuildError(f"{table}: profiling exceeds max_rows={max_rows}; no uniqueness is inferred from a sample.")
    return profile(identifier(spec.get("name", table)), spec, rows)
